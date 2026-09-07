"""Converte os relatorios de Markdown para .docx.

Uso:
    python experimento/gerar_docx.py                  # todos os .md de relatorios/
    python experimento/gerar_docx.py caminho/arq.md   # um arquivo especifico

O .md continua sendo a fonte da verdade. O .docx e' derivado e regenerado -- nunca
edite o .docx diretamente, porque a proxima execucao o sobrescreve.

Subconjunto de Markdown suportado: titulos (# a ####), paragrafos, negrito, italico,
codigo inline, listas com marcador e numeradas, tabelas com barra vertical, imagens,
citacoes e blocos de codigo cercados.
"""
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RELATORIOS = os.path.join(RAIZ, "relatorios")

LARGURA_UTIL_CM = 16.0          # A4 com margens de 2,5 cm
RE_INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)")
RE_IMAGEM = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)\s*$")
RE_TITULO = re.compile(r"^(?P<nivel>#{1,6})\s+(?P<texto>.+?)\s*#*\s*$")
RE_LISTA_NUM = re.compile(r"^\s*\d+[.)]\s+(?P<texto>.+)$")
RE_LISTA_MARC = re.compile(r"^\s*[-*+]\s+(?P<texto>.+)$")
RE_SEPARADOR = re.compile(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$")


def escrever_inline(paragrafo, texto):
    """Quebra o texto em trechos formatados e os adiciona como runs."""
    for trecho in RE_INLINE.split(texto):
        if not trecho:
            continue
        if trecho.startswith("**") and trecho.endswith("**"):
            paragrafo.add_run(trecho[2:-2]).bold = True
        elif trecho.startswith("`") and trecho.endswith("`"):
            r = paragrafo.add_run(trecho[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
        elif trecho.startswith("*") and trecho.endswith("*"):
            paragrafo.add_run(trecho[1:-1]).italic = True
        else:
            paragrafo.add_run(trecho)


def linha_de_tabela(linha):
    return linha.lstrip().startswith("|") and linha.rstrip().endswith("|")


def celulas(linha):
    return [c.strip() for c in linha.strip().strip("|").split("|")]


def eh_separador_de_tabela(linha):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", linha.strip()))


def inserir_tabela(doc, blocos):
    """blocos: linhas de uma tabela markdown, ja sem a linha separadora."""
    dados = [celulas(l) for l in blocos]
    n_col = max(len(l) for l in dados)
    tabela = doc.add_table(rows=len(dados), cols=n_col)
    tabela.style = "Table Grid"
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, linha in enumerate(dados):
        for j in range(n_col):
            cel = tabela.cell(i, j)
            cel.text = ""
            p = cel.paragraphs[0]
            escrever_inline(p, linha[j] if j < len(linha) else "")
            for run in p.runs:
                run.font.size = Pt(9)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def inserir_imagem(doc, src, alt, base):
    for cand in (os.path.join(base, src), os.path.join(RAIZ, src), src):
        caminho = os.path.normpath(cand)
        if os.path.isfile(caminho):
            doc.add_picture(caminho, width=Cm(LARGURA_UTIL_CM))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if alt:
                leg = doc.add_paragraph()
                leg.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = leg.add_run(alt)
                r.italic = True
                r.font.size = Pt(9)
            return True
    # Imagem ausente e' erro visivel, nao silencioso: o .docx vai para a banca.
    p = doc.add_paragraph()
    p.add_run("[IMAGEM NAO ENCONTRADA: " + src + "]").bold = True
    print("   AVISO: imagem nao encontrada -> " + src)
    return False


def converter(caminho_md):
    base = os.path.dirname(os.path.abspath(caminho_md))
    with open(caminho_md, encoding="utf-8") as fh:
        linhas = fh.read().splitlines()

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = s.right_margin = Cm(2.5)

    i, em_codigo, buffer_codigo = 0, False, []
    while i < len(linhas):
        linha = linhas[i]

        if linha.strip().startswith("```"):
            if em_codigo:
                p = doc.add_paragraph()
                r = p.add_run("\n".join(buffer_codigo))
                r.font.name = "Consolas"
                r.font.size = Pt(9)
                buffer_codigo, em_codigo = [], False
            else:
                em_codigo = True
            i += 1
            continue
        if em_codigo:
            buffer_codigo.append(linha)
            i += 1
            continue

        if not linha.strip() or RE_SEPARADOR.match(linha):
            i += 1
            continue

        if linha_de_tabela(linha):
            bloco = []
            while i < len(linhas) and linha_de_tabela(linhas[i]):
                if not eh_separador_de_tabela(linhas[i]):
                    bloco.append(linhas[i])
                i += 1
            if bloco:
                inserir_tabela(doc, bloco)
            continue

        m = RE_IMAGEM.match(linha.strip())
        if m:
            inserir_imagem(doc, m.group("src"), m.group("alt"), base)
            i += 1
            continue

        m = RE_TITULO.match(linha)
        if m:
            doc.add_heading(m.group("texto"), level=min(len(m.group("nivel")), 4))
            i += 1
            continue

        m = RE_LISTA_NUM.match(linha)
        if m:
            escrever_inline(doc.add_paragraph(style="List Number"), m.group("texto"))
            i += 1
            continue

        m = RE_LISTA_MARC.match(linha)
        if m:
            escrever_inline(doc.add_paragraph(style="List Bullet"), m.group("texto"))
            i += 1
            continue

        if linha.lstrip().startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            escrever_inline(p, linha.lstrip().lstrip(">").strip())
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        # Paragrafo: junta linhas ate a proxima linha em branco ou marcador de bloco.
        bloco = []
        while i < len(linhas) and linhas[i].strip() and not (
            RE_TITULO.match(linhas[i]) or linha_de_tabela(linhas[i])
            or RE_LISTA_NUM.match(linhas[i]) or RE_LISTA_MARC.match(linhas[i])
            or RE_SEPARADOR.match(linhas[i]) or linhas[i].strip().startswith("```")
            or linhas[i].lstrip().startswith(">") or RE_IMAGEM.match(linhas[i].strip())
        ):
            bloco.append(linhas[i].strip())
            i += 1
        if bloco:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            escrever_inline(p, " ".join(bloco))

    saida = os.path.splitext(caminho_md)[0] + ".docx"
    doc.save(saida)
    return saida


def main():
    alvos = sys.argv[1:]
    if not alvos:
        if not os.path.isdir(RELATORIOS):
            print("ERRO: " + RELATORIOS + " nao existe.")
            return 1
        alvos = sorted(os.path.join(RELATORIOS, f) for f in os.listdir(RELATORIOS)
                       if f.lower().endswith(".md"))
    if not alvos:
        print("Nenhum .md encontrado em relatorios/.")
        return 1
    for md in alvos:
        if not os.path.isfile(md):
            print("ERRO: nao encontrado -> " + md)
            return 1
        saida = converter(md)
        kb = os.path.getsize(saida) / 1024
        print("-> " + os.path.relpath(saida, RAIZ) + "  ({:.0f} KB)".format(kb))
    return 0


if __name__ == "__main__":
    sys.exit(main())
